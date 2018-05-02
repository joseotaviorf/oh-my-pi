# -*- coding: utf-8 -*-
# openpyxl==2.4.8
import re
import sys
from datetime import datetime

import boto3
import pandas as pd
from openpyxl import load_workbook
from qa_python_utils.default_logger import logger, _logger

from bietlejuice.jobs.base.base_etl import BaseETL

args = sys.argv


class AgentsFinanceData(object):
    @logger
    def __init__(self, file_name=None):
        self.s3_client = boto3.client('s3')
        self.file_name = file_name

        if file_name:
            self.workbook = load_workbook(file_name)

    @logger
    def get_commissions_data(self):
        return None

    @logger
    def get_contracts_data(self):
        return None

    @logger
    def _get_commission_data_2017_until_aug_old(self):
        """
            Method used to reprocess old commission data (from Jan 2017 to Aug 2017)
            *Warning*: delete old data before executing this method, otherwise duplicated data will appear
        """
        summary = self.workbook.get_sheet_names()[0]
        summary_worksheet = self.workbook.get_sheet_by_name(summary)

        # FIXME: fix s3 read
        file_obj = self.s3_client.get_object(
            Bucket='5a-datalake',
            Key='raw/files/agent_commissions/Horas e comissões - Consolidado Corretores-2017.V2.xlsm')

        df_summary_values = pd.read_excel(file_obj['Body'], sheetname='Resumo')
        df_commissions = df_summary_values[df_summary_values.Item == u'4 Comissão']

        dates = []
        df_agents = pd.DataFrame(columns=['agent_id', 'agent_name', 'dt', 'percentage', 'value'])
        for row in summary_worksheet.iter_rows():
            if row[0].row == 1:
                i = 3
                while isinstance(row[i].value, datetime):
                    dates.append(row[i].value.date())
                    i += 1
            else:
                agent_id = row[0].value
                agent_name = row[1].value

                if row[2].value and 'Comissão' in row[2].value:
                    i = 3
                    percentage = 0
                    while row[i].value is not None:
                        regex_result = re.search('\*(.*)%', row[i].value)
                        if regex_result:
                            percentage = int(regex_result.group(1)) / 100.
                        dt = dates[i - 3]
                        df_agents = df_agents.append([
                            {
                                'agent_id': agent_id,
                                'agent_name': agent_name,
                                'dt': dt,
                                'percentage': round(percentage, 2),
                                'value': df_commissions[df_commissions['ID Corretor'] == agent_id].iloc[0][i]
                            }
                        ])
                        i += 1

        return df_agents

    @logger
    def _get_commission_data_2016_old(self):
        """
            Method used to reprocess old commission data (from 2015 to 2016)
            *Warning*: delete old data before executing this method, otherwise duplicated data will appear
        """
        summary = self.workbook.get_sheet_names()[0]
        summary_worksheet = self.workbook.get_sheet_by_name(summary)

        dates = []
        df_agents = pd.DataFrame(columns=['agent_id', 'agent_name', 'dt', 'percentage', 'value'])
        for row in summary_worksheet.iter_rows():
            if row[0].row == 1:
                i = 2
                while isinstance(row[i].value, datetime):
                    dates.append(row[i].value.date())
                    i += 1
            else:
                agent_name = row[0].value
                if row[1].value and 'Comissão' in row[1].value \
                        and agent_name is not None and re.search('=[A-Z]\d*', agent_name) is None:
                    i = 2
                    percentage = 0
                    while i - 2 < len(dates):
                        if row[i].value is None:
                            i += 1
                            continue

                        regex_result = re.search('\*((.|,)*)%?', row[i].value)
                        if regex_result:
                            percentage = float(regex_result.group(1))

                            # workaround for CLT agents
                            if percentage == 0.05:
                                percentage = 0.5

                        dt = dates[i - 2]
                        df_agents = df_agents.append([
                            {
                                'agent_id': None,
                                'agent_name': agent_name,
                                'dt': dt,
                                'percentage': round(percentage, 2),
                                'value': None
                            }
                        ])
                        i += 1

        return df_agents

    @logger
    def _get_contracts_data_2017_until_aug_old(self):
        """
            Method used to reprocess old contract data (from Jan 2017 to Aug 2017)
            *Warning*: delete old data before executing this method, otherwise duplicated data will appear
        """
        # FIXME: fix s3 read
        file_obj = self.s3_client.get_object(
            Bucket='5a-datalake',
            Key='raw/files/agent_commissions/Horas e comissões - Consolidado Corretores-2017.V2.xlsm')

        df_contract_values = pd.read_excel(file_obj['Body'], sheetname='Contratos')
        df_contract_values = df_contract_values.astype(object).where(pd.notnull(df_contract_values), None)

        df_contracts = pd.DataFrame(
            columns=['agent_id', 'agent_name', 'status', 'property_id', 'contract_id', 'signature_date', 'rent'])
        for index, row in df_contract_values.iterrows():
            agent_id = row[0] if isinstance(row[0], int) else None
            agent_name = row[1]
            status = row[2]
            property_id = contract_id = None
            if row[3]:
                if len(str(int(row[3]))) == 5:
                    property_id = int(row[3])
                if len(str(int(row[3]))) == 4:
                    contract_id = int(row[3])

            signature_date = row[5]
            rent = row[6]

            df_contracts = df_contracts.append([
                {
                    'agent_id': agent_id,
                    'agent_name': agent_name,
                    'status': status,
                    'property_id': property_id,
                    'contract_id': contract_id,
                    'signature_date': signature_date,
                    'rent': rent
                }
            ])

        return df_contracts.astype(object).where(pd.notnull(df_contracts), None)

    @logger
    def _get_contracts_data_2016_old(self):
        """
            Method used to reprocess old contract data (from 2015 to 2016)
            *Warning*: delete old data before executing this method, otherwise duplicated data will appear
        """
        # FIXME: fix s3 read
        file_obj = self.s3_client.get_object(
            Bucket='5a-datalake',
            Key='raw/files/agent_commissions/Horas e comissões - Consolidado Corretores.xlsm')

        df_contract_values = pd.read_excel(file_obj['Body'], sheetname='Contratos')
        df_contract_values = df_contract_values.astype(object).where(pd.notnull(df_contract_values), None)

        df_contracts = pd.DataFrame(
            columns=['agent_id', 'agent_name', 'status', 'property_id', 'contract_id', 'signature_date', 'rent'])
        for index, row in df_contract_values.iterrows():
            agent_name = row[0]
            status = row[1]
            property_id = int(row[2]) if isinstance(row[2], float) else None
            signature_date = row[4]
            rent = row[5]

            df_contracts = df_contracts.append([
                {
                    'agent_id': None,
                    'agent_name': agent_name,
                    'status': status,
                    'property_id': property_id,
                    'contract_id': None,
                    'signature_date': signature_date,
                    'rent': rent
                }
            ])

        return df_contracts.astype(object).where(pd.notnull(df_contracts), None)

    @logger(exclude='df')
    def save_df_to_ods(self, df, table_name):
        if df is None or df.empty:
            _logger.warn('m=save_df_to_ods, msg=dataframe is empty!')
            return

        BaseETL.dataframe_to_ods(df=df, table_name=table_name, append=True)

    @logger
    def delete_from_aug_2017(self):
        # TODO
        pass


if __name__ == '__main__':
    # TODO: get file_name from airflow env var
    file_name = None
    agents_data = AgentsFinanceData(file_name)
    agents_data.delete_from_aug_2017()

    if args[1] == 'agents_commission':
        df = agents_data.get_commissions_data()
        agents_data.save_df_to_ods(df, 'files.finance_agents_commission')
    elif args[1] == 'agents_contracts':
        df = agents_data.get_contracts_data()
        agents_data.save_df_to_ods(df, 'files.finance_agents_contract')
    else:
        _logger.info('m=__main__, msg=arg \'{}\' not recognized'.format(args[1]))
